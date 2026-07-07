import csv
import hashlib
import io
import re

from app import config
from app.db import queries

FREE_EMAIL_DOMAINS = {
    "gmail.com", "googlemail.com", "outlook.com", "hotmail.com", "live.com", "msn.com",
    "yahoo.com", "ymail.com", "aol.com", "icloud.com", "me.com", "mac.com",
    "proton.me", "protonmail.com", "gmx.com", "gmx.de", "gmx.net", "mail.com",
    "web.de", "t-online.de", "freenet.de", "orange.fr", "wanadoo.fr", "free.fr",
    "laposte.net", "yandex.ru", "yandex.com", "mail.ru", "zoho.com", "fastmail.com",
}

_COLUMN_MAP = {
    "company_name": ["company", "company name", "organization", "account", "companyname",
                     "company name for emails", "account name", "organization name"],
    "domain": ["domain", "website", "url", "company domain", "website url",
               "company website", "organization website url", "site"],
    "contact_name": ["name", "full name", "contact", "contact name", "person name",
                     "lead name"],
    "contact_title": ["title", "job title", "position", "role", "seniority title"],
    "email": ["email", "email address", "work email", "person email", "verified email",
              "primary email"],
    "linkedin_url": ["linkedin", "linkedin url", "linkedin profile", "person linkedin url",
                     "profile url", "linkedin profile url"],
    "_first_name": ["first name", "firstname", "first"],
    "_last_name": ["last name", "lastname", "last", "surname"],
}

SHEET_STUB_ROWS = [
    {"Company": "Acme Logistics", "Website": "acme-logistics.com",
     "First Name": "Sam", "Last Name": "Lee", "Email": "sam@acme-logistics.com"},
    {"Company": "Beta Freight", "Website": "betafreight.io",
     "First Name": "Jordan", "Last Name": "Kim", "Email": "jordan@betafreight.io"},
]


def parse_csv(raw_bytes):
    text = raw_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return [], ["File has no header row."]
    return _rows_from_dicts(list(reader), reader.fieldnames)


def parse_xlsx(raw_bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception:
        return [], ["Could not read the Excel file. Save it as .xlsx and retry."]
    sheet = wb.worksheets[0]
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header or not any(header):
        return [], ["File has no header row."]
    fieldnames = [str(h).strip() if h is not None else "" for h in header]
    dict_rows = []
    for values in rows_iter:
        if values is None or not any(v is not None and str(v).strip() for v in values):
            continue
        dict_rows.append({
            fieldnames[i]: ("" if v is None else str(v).strip())
            for i, v in enumerate(values) if i < len(fieldnames) and fieldnames[i]
        })
    wb.close()
    return _rows_from_dicts(dict_rows, fieldnames)


def parse_google_sheet(url_or_id):
    sheet_id = _sheet_id(url_or_id)
    if not sheet_id:
        return [], ["That doesn't look like a Google Sheets link or ID."]
    if config.USE_FETCH_STUB:
        fieldnames = list(SHEET_STUB_ROWS[0].keys())
        return _rows_from_dicts([dict(r) for r in SHEET_STUB_ROWS], fieldnames)
    raw = _fetch_sheet_csv(sheet_id)
    if raw is None:
        return [], [
            "Could not read the sheet. Share it as 'Anyone with the link — Viewer' and retry."
        ]
    return parse_csv(raw)


def _sheet_id(value):
    value = (value or "").strip()
    m = re.search(r"/d/([A-Za-z0-9_-]{20,})", value)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{20,}", value):
        return value
    return ""


def _fetch_sheet_csv(sheet_id):
    import httpx

    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
    try:
        resp = httpx.get(url, follow_redirects=True, timeout=20)
    except httpx.HTTPError:
        return None
    if resp.status_code >= 400 or b"<html" in resp.content[:200].lower():
        return None
    return resp.content


def _rows_from_dicts(dict_rows, fieldnames):
    header_lookup = _build_header_lookup(fieldnames)
    rows = []
    errors = []
    for i, raw in enumerate(dict_rows, start=2):
        lead = _map_row(raw, header_lookup)
        if not lead["company_name"] and not lead["domain"] and not lead["email"]:
            errors.append(f"Row {i}: no company, domain, or email.")
            continue
        lead["domain"] = _clean_domain(lead["domain"] or _domain_from_email(lead["email"]))
        lead["dedupe_hash"] = _dedupe_hash(lead)
        rows.append(lead)
    return rows, errors


def _build_header_lookup(fieldnames):
    lookup = {}
    for field in fieldnames:
        key = field.strip().lower().replace("_", " ")
        for target, aliases in _COLUMN_MAP.items():
            if target in lookup:
                continue
            if key == target or key in aliases:
                lookup[target] = field
                break
    return lookup


def _map_row(raw, header_lookup):
    def val(target):
        col = header_lookup.get(target)
        return (raw.get(col) or "").strip() if col else ""

    contact_name = val("contact_name")
    if not contact_name:
        contact_name = " ".join(p for p in [val("_first_name"), val("_last_name")] if p)

    return {
        "company_name": val("company_name"),
        "domain": val("domain"),
        "contact_name": contact_name,
        "contact_title": val("contact_title"),
        "email": val("email").lower(),
        "linkedin_url": val("linkedin_url"),
        "raw": {k: v for k, v in raw.items() if k},
    }


def _clean_domain(value):
    if not value:
        return ""
    value = re.sub(r"^https?://", "", value.strip().lower())
    value = value.split("/")[0]
    domain = value[4:] if value.startswith("www.") else value
    return "" if domain in FREE_EMAIL_DOMAINS else domain


def _domain_from_email(email):
    domain = email.split("@")[1] if "@" in email else ""
    return "" if domain in FREE_EMAIL_DOMAINS else domain


def _dedupe_hash(lead):
    key = (lead["email"] or lead["domain"] or lead["company_name"]).lower()
    return hashlib.sha256(key.encode("utf-8")).hexdigest()[:16]


def import_rows(list_id, rows):
    seen = queries.existing_dedupe_hashes(list_id)
    inserted = 0
    duplicates = 0
    for lead in rows:
        if lead["dedupe_hash"] in seen:
            duplicates += 1
            continue
        seen.add(lead["dedupe_hash"])
        queries.insert_lead(list_id, lead)
        inserted += 1
    return inserted, duplicates
